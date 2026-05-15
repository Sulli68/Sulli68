"""
generate_pptx_v2.py
High-fidelity testachats/testaankoop brand presentation – 5 slides.
"""

from pptx import Presentation
from pptx.util import Inches, Pt, Emu
from pptx.dml.color import RGBColor
from pptx.enum.text import PP_ALIGN, MSO_ANCHOR
from pptx.util import Inches, Pt
from pptx.oxml.ns import qn
from lxml import etree
import copy

# ── colour palette ────────────────────────────────────────────────────────────
RED    = RGBColor(0xE8, 0x00, 0x3D)
ORANGE = RGBColor(0xF0, 0x7D, 0x00)
AMBER  = RGBColor(0xF5, 0xA6, 0x23)
DARK   = RGBColor(0x22, 0x22, 0x22)
GRAY   = RGBColor(0x77, 0x77, 0x77)
LGRAY  = RGBColor(0xF3, 0xF3, 0xF3)
MGRAY  = RGBColor(0xCE, 0xCE, 0xCE)
WHITE  = RGBColor(0xFF, 0xFF, 0xFF)
PINK_L = RGBColor(0xFF, 0xE5, 0xEE)
ORAN_L = RGBColor(0xFE, 0xF0, 0xE2)
YELL_L = RGBColor(0xFF, 0xF8, 0xE5)
GOLD   = RGBColor(0x7A, 0x5C, 0x00)
OPPORT = RGBColor(0xFF, 0xCC, 0xD8)

FONT = 'Calibri'

# ── helpers ───────────────────────────────────────────────────────────────────

def _in(val):
    return Inches(val)

def _set_run(run, size, bold, italic, color, name=FONT):
    run.font.name  = name
    run.font.size  = Pt(size)
    run.font.bold  = bold
    run.font.italic = italic
    run.font.color.rgb = color

def _no_line(shp):
    shp.line.fill.background()

def rect(slide, l, t, w, h, fill, line=None, lw=0.75):
    from pptx.util import Pt as _Pt
    shp = slide.shapes.add_shape(1, _in(l), _in(t), _in(w), _in(h))
    shp.fill.solid()
    shp.fill.fore_color.rgb = fill
    if line:
        shp.line.color.rgb = line
        shp.line.width = _Pt(lw)
    else:
        _no_line(shp)
    return shp

def oval(slide, l, t, w, h, fill):
    shp = slide.shapes.add_shape(9, _in(l), _in(t), _in(w), _in(h))
    shp.fill.solid()
    shp.fill.fore_color.rgb = fill
    _no_line(shp)
    return shp

def shape_txt(slide, text, l, t, w, h, fill, size, bold, color, align, v_anchor=MSO_ANCHOR.MIDDLE):
    shp = slide.shapes.add_shape(1, _in(l), _in(t), _in(w), _in(h))
    shp.fill.solid()
    shp.fill.fore_color.rgb = fill
    _no_line(shp)
    tf = shp.text_frame
    tf.word_wrap = True
    tf.vertical_anchor = v_anchor
    lines = text.split('\n')
    for i, line in enumerate(lines):
        if i == 0:
            p = tf.paragraphs[0]
        else:
            p = tf.add_paragraph()
        p.alignment = align
        run = p.add_run()
        run.text = line
        _set_run(run, size, bold, False, color)
    return shp

def txt(slide, text, l, t, w, h, size, bold, italic, color, align):
    txb = slide.shapes.add_textbox(_in(l), _in(t), _in(w), _in(h))
    tf  = txb.text_frame
    tf.word_wrap = True
    lines = text.split('\n')
    for i, line in enumerate(lines):
        if i == 0:
            p = tf.paragraphs[0]
        else:
            p = tf.add_paragraph()
        p.alignment = align
        run = p.add_run()
        run.text = line
        _set_run(run, size, bold, italic, color)
    return txb

def mtxt(slide, paras, l, t, w, h):
    txb = slide.shapes.add_textbox(_in(l), _in(t), _in(w), _in(h))
    tf  = txb.text_frame
    tf.word_wrap = True
    for i, pd in enumerate(paras):
        if i == 0:
            p = tf.paragraphs[0]
        else:
            p = tf.add_paragraph()
        p.alignment = pd.get('align', PP_ALIGN.LEFT)
        if pd.get('space_before'):
            p.space_before = Pt(pd['space_before'])
        run = p.add_run()
        run.text = pd['text']
        _set_run(run, pd['size'], pd['bold'], pd.get('italic', False), pd['color'])
    return txb

def badge(slide, num, cx, cy):
    r = 0.22
    o = oval(slide, cx - r, cy - r, r*2, r*2, RED)
    tf = o.text_frame
    tf.vertical_anchor = MSO_ANCHOR.MIDDLE
    p = tf.paragraphs[0]
    p.alignment = PP_ALIGN.CENTER
    run = p.add_run()
    run.text = str(num)
    _set_run(run, 9, True, False, WHITE)

def divider(slide, l, t, w):
    rect(slide, l, t, w, 0.04, RED)

def logo(slide):
    txb = slide.shapes.add_textbox(_in(0.28), _in(7.1), _in(2.0), _in(0.32))
    tf  = txb.text_frame
    tf.word_wrap = False
    p = tf.paragraphs[0]
    p.alignment = PP_ALIGN.LEFT
    r1 = p.add_run(); r1.text = 'testachats'
    _set_run(r1, 8, True, False, RED)
    r2 = p.add_run(); r2.text = '  '
    _set_run(r2, 8, False, False, DARK)
    r3 = p.add_run(); r3.text = 'testaankoop'
    _set_run(r3, 8, True, False, DARK)

def pnum(slide, n):
    txt(slide, str(n), 12.9, 7.15, 0.35, 0.28, 7, False, False, GRAY, PP_ALIGN.RIGHT)

def icon_oval(slide, symbol, cx, cy, r, bg, sym_color, sym_size):
    o = oval(slide, cx - r, cy - r, r*2, r*2, bg)
    tf = o.text_frame
    tf.vertical_anchor = MSO_ANCHOR.MIDDLE
    p = tf.paragraphs[0]
    p.alignment = PP_ALIGN.CENTER
    run = p.add_run()
    run.text = symbol
    _set_run(run, sym_size, True, False, sym_color)
    return o

def callout(slide, text, l, t, w, h, bg, border_col, text_col, size):
    bw = 0.06
    # left border bar
    rect(slide, l, t, bw, h, border_col)
    # background rect
    rect(slide, l + bw, t, w - bw, h, bg)
    # text box inside
    pad = 0.08
    tb = slide.shapes.add_textbox(_in(l + bw + pad), _in(t + pad),
                                   _in(w - bw - 2*pad), _in(h - 2*pad))
    tf = tb.text_frame
    tf.word_wrap = True
    lines = text.split('\n')
    for i, line in enumerate(lines):
        if i == 0:
            p = tf.paragraphs[0]
        else:
            p = tf.add_paragraph()
        p.alignment = PP_ALIGN.LEFT
        run = p.add_run()
        run.text = line
        _set_run(run, size, False, True, text_col)

# ── build presentation ────────────────────────────────────────────────────────

prs = Presentation()
prs.slide_width  = Inches(13.33)
prs.slide_height = Inches(7.5)

blank_layout = prs.slide_layouts[6]   # completely blank

# ─────────────────────────────────────────────────────────────────────────────
# SLIDE 1 – "WHAT IS TRULY ESSENTIAL FOR CONSUMERS?"
# ─────────────────────────────────────────────────────────────────────────────
sl = prs.slides.add_slide(blank_layout)

# white bg
rect(sl, 0, 0, 13.33, 7.5, WHITE)

# badge
badge(sl, 1, 0.52, 0.45)

# title
mtxt(sl, [
    {'text': 'WHAT IS TRULY ESSENTIAL', 'size': 22, 'bold': True, 'color': DARK,  'align': PP_ALIGN.LEFT},
    {'text': 'FOR CONSUMERS?',          'size': 22, 'bold': True, 'color': RED,   'align': PP_ALIGN.LEFT},
], 0.9, 0.1, 10.5, 1.0)

divider(sl, 0.9, 1.12, 0.9)

txt(sl, 'Our value is built on trust, expertise and impact.',
    0.28, 1.2, 11.5, 0.32, 10, False, True, GRAY, PP_ALIGN.LEFT)

# 4 column cards
MARGIN = 0.28
GAP    = 0.12
CW4    = 3.0
xs     = [0.28, 3.40, 6.52, 9.64]
y_top  = 1.6
card_h = 3.85

cols = [
    {
        'sym': '⚖', 'bg_sym': PINK_L, 'sym_col': RED,
        'title': 'TRUST &\nINDEPENDENCE', 'tcol': DARK,
        'bullets': ['• Independent voice', '• High credibility', '• Consumer-first mission'],
    },
    {
        'sym': '⚙', 'bg_sym': PINK_L, 'sym_col': ORANGE,
        'title': 'PRACTICAL CONSUMER\nSUPPORT', 'tcol': ORANGE,
        'bullets': ['• Rights & disputes', '• Financial guidance', '• Energy', '• Major purchases', '• Contracts'],
    },
    {
        'sym': '◈', 'bg_sym': PINK_L, 'sym_col': ORANGE,
        'title': 'EXPERTISE &\nDECISION SUPPORT', 'tcol': ORANGE,
        'bullets': ['• Testing capability', '• Structured data', '• Comparison intelligence', '• Help consumers make decisions'],
    },
    {
        'sym': '◉', 'bg_sym': PINK_L, 'sym_col': ORANGE,
        'title': 'SIMPLICITY &\nGUIDANCE', 'tcol': ORANGE,
        'bullets': ['• Clear answers', '• Actionable advice', '• Save time & money', '• Build confidence'],
    },
]

for i, col in enumerate(cols):
    x = xs[i]
    # card bg
    rect(sl, x, y_top, CW4, card_h, LGRAY)
    # icon oval
    icon_oval(sl, col['sym'], x + CW4/2, y_top + 0.48, 0.3, col['bg_sym'], col['sym_col'], 14)
    # title
    shape_txt(sl, col['title'], x + 0.05, y_top + 0.88, 2.9, 0.55,
              LGRAY, 8.5, True, col['tcol'], PP_ALIGN.CENTER)
    # separator
    rect(sl, x + 0.12, y_top + 1.5, 2.76, 0.03, MGRAY)
    # bullets
    bul_y = y_top + 1.6
    for b in col['bullets']:
        txt(sl, b, x + 0.15, bul_y, 2.7, 0.3, 8, False, False, DARK, PP_ALIGN.LEFT)
        bul_y += 0.3

# bottom stats
# left LGRAY box
rect(sl, 0.28, 5.55, 2.3, 1.68, LGRAY)
txt(sl, '96%',              0.28, 5.6,  2.3, 0.58, 38, True,  False, DARK, PP_ALIGN.CENTER)
txt(sl, 'aided awareness',  0.28, 6.22, 2.3, 0.28, 7.5, False, False, GRAY, PP_ALIGN.CENTER)
txt(sl, '(Brand Tracker Q4 2024)', 0.28, 6.52, 2.3, 0.28, 7, False, True, GRAY, PP_ALIGN.CENTER)

# middle stats
txt(sl, 'Top relevance scores:', 2.72, 5.55, 2.9, 0.28, 8, True, False, ORANGE, PP_ALIGN.LEFT)
stats_items = ['• 78% fiscal questions', '• 76% health questions', '• 72% property decisions']
sy = 5.85
for s in stats_items:
    txt(sl, s, 2.72, sy, 2.9, 0.26, 8, False, False, DARK, PP_ALIGN.LEFT)
    sy += 0.26
txt(sl, '(Brand Tracker Q4 2024)', 2.72, sy, 2.9, 0.26, 7, False, True, GRAY, PP_ALIGN.LEFT)

txt(sl, 'Top traffic hubs:', 5.75, 5.55, 2.9, 0.28, 8, True, False, ORANGE, PP_ALIGN.LEFT)
hub_items = ['• 20% Personal Finance', '• 16.8% Home Equipment', '• 16.8% High Tech']
hy = 5.85
for h in hub_items:
    txt(sl, h, 5.75, hy, 2.9, 0.26, 8, False, False, DARK, PP_ALIGN.LEFT)
    hy += 0.26
txt(sl, '(Hub Online Sessions)', 5.75, hy, 2.9, 0.26, 7, False, True, GRAY, PP_ALIGN.LEFT)

callout(sl,
        'Relevance is the #1\ndriver of consideration\n(43%)\n(Brand Tracker Q4 2024)',
        8.78, 5.5, 4.28, 1.75, PINK_L, RED, RED, 8.5)

logo(sl); pnum(sl, 1)


# ─────────────────────────────────────────────────────────────────────────────
# SLIDE 2 – "WHERE SHOULD WE PRIORITISE, SIMPLIFY OR STOP?"
# ─────────────────────────────────────────────────────────────────────────────
sl = prs.slides.add_slide(blank_layout)
rect(sl, 0, 0, 13.33, 7.5, WHITE)
badge(sl, 2, 0.52, 0.45)

mtxt(sl, [
    {'text': 'WHERE SHOULD WE PRIORITISE,', 'size': 21, 'bold': True, 'color': DARK, 'align': PP_ALIGN.LEFT},
    {'text': 'SIMPLIFY OR STOP?',           'size': 21, 'bold': True, 'color': RED,  'align': PP_ALIGN.LEFT},
], 0.9, 0.1, 11.5, 0.98)

divider(sl, 0.9, 1.1, 0.9)
txt(sl, 'Focus on what creates the most value. Simplify to increase impact.',
    0.28, 1.18, 12.8, 0.32, 10, False, True, GRAY, PP_ALIGN.CENTER)

CW3 = 4.17
xs3 = [0.28, 4.57, 8.86]

# column headers
shape_txt(sl, 'PRIORITISE MORE',    0.28, 1.58, CW3, 0.42, RED,   11, True, WHITE, PP_ALIGN.CENTER)
shape_txt(sl, 'SIMPLIFY',           4.57, 1.58, CW3, 0.42, ORANGE,11, True, WHITE, PP_ALIGN.CENTER)
shape_txt(sl, 'REDUCE / REASSESS',  8.86, 1.58, CW3, 0.42, AMBER, 11, True, WHITE, PP_ALIGN.CENTER)

# sub-headers
sub_hdrs = ['Invest in high-value domains', 'Reduce operational complexity', 'Rationalise lower impact areas']
for i, sh in enumerate(sub_hdrs):
    txt(sl, sh, xs3[i], 2.05, CW3, 0.28, 8.5, True, False, DARK, PP_ALIGN.LEFT)

# items with colored dots
col1_items = ['Personal Finance', 'Rights & Protection', 'Housing & Energy', 'High Tech', 'High-impact consumer decisions']
col2_items = ['Shared hub-based workflows', 'Reduce duplicate layers', 'Simplify web/print production', 'AI-assisted repetitive tasks', 'Faster multi-channel publishing']
col3_items = ['Low-interest niche testing', 'Fragmented micro-topics', 'Duplicate workflows', 'Separation between publication brands', 'Heavy manual adaptation processes']
dot_cols = [RED, ORANGE, AMBER]
all_cols = [col1_items, col2_items, col3_items]

for ci, (col_items, dot_col) in enumerate(zip(all_cols, dot_cols)):
    x = xs3[ci]
    iy = 2.38
    for item in col_items:
        oval(sl, x + 0.0, iy + 0.06, 0.18, 0.18, dot_col)
        txt(sl, item, x + 0.24, iy, CW3 - 0.28, 0.3, 8.5, False, False, DARK, PP_ALIGN.LEFT)
        iy += 0.36

# callout boxes
callout(sl, '54% of online sessions concentrated in Personal Finance, Home Equipment and High Tech. (Hub Online Sessions)',
        0.28, 4.25, 4.17, 0.92, PINK_L, RED, RED, 7.5)
callout(sl, 'Major opportunity to reduce complexity and speed up publishing. (Internal observation)',
        4.57, 4.25, 4.17, 0.92, ORAN_L, ORANGE, ORANGE, 7.5)
callout(sl, 'Lower traffic hubs:\nShopping 3.1% | Rights 5.4% | Mobility 5.5%\n(Hub Online Sessions)',
        8.86, 4.25, 4.17, 0.92, YELL_L, AMBER, GOLD, 7.5)

# bottom AI / Human section
rect(sl, 0, 5.25, 13.33, 0.03, MGRAY)
txt(sl, 'AI SHOULD SUPPORT', 0.28, 5.32, 5.8, 0.3, 8.5, True, False, DARK, PP_ALIGN.LEFT)
ai_items = ['• First drafts', '• Translation & SEO', '• Workflow acceleration']
aiy = 5.65
for ai in ai_items:
    txt(sl, ai, 0.28, aiy, 5.8, 0.28, 8.5, False, False, DARK, PP_ALIGN.LEFT)
    aiy += 0.3

rect(sl, 6.67, 5.22, 0.03, 1.8, MGRAY)
txt(sl, 'HUMAN EXPERTISE REMAINS ESSENTIAL FOR', 6.87, 5.32, 6.0, 0.3, 8.5, True, False, DARK, PP_ALIGN.LEFT)
hu_items = ['• Legal interpretation', '• Validation & credibility', '• Consumer advocacy']
huy = 5.65
for hu in hu_items:
    txt(sl, hu, 6.87, huy, 6.0, 0.28, 8.5, False, False, DARK, PP_ALIGN.LEFT)
    huy += 0.3

logo(sl); pnum(sl, 2)


# ─────────────────────────────────────────────────────────────────────────────
# SLIDE 3 – "OPERATIONAL CONSEQUENCES & FUTURE ORGANISATION"
# ─────────────────────────────────────────────────────────────────────────────
sl = prs.slides.add_slide(blank_layout)
rect(sl, 0, 0, 13.33, 7.5, WHITE)
badge(sl, 3, 0.52, 0.45)

mtxt(sl, [
    {'text': 'OPERATIONAL CONSEQUENCES', 'size': 21, 'bold': True, 'color': DARK, 'align': PP_ALIGN.LEFT},
    {'text': '& FUTURE ORGANISATION',    'size': 21, 'bold': True, 'color': RED,  'align': PP_ALIGN.LEFT},
], 0.9, 0.1, 11.5, 0.98)

divider(sl, 0.9, 1.1, 0.9)
txt(sl, 'From publication silos to integrated consumer hubs.',
    0.28, 1.18, 12.0, 0.32, 10, False, True, GRAY, PP_ALIGN.LEFT)

# TODAY column
shape_txt(sl, 'TODAY', 0.28, 1.58, 5.8, 0.4, LGRAY, 11, True, DARK, PP_ALIGN.CENTER)
txt(sl, 'Publication-centric & fragmented', 0.28, 2.02, 5.8, 0.28, 8.5, False, True, GRAY, PP_ALIGN.CENTER)

# 4 publication ovals
pub_labels = ['TA', 'BD', 'TS', 'TAI']
pub_xs = [0.38, 1.58, 2.78, 3.98]
for label, px in zip(pub_labels, pub_xs):
    o = oval(sl, px, 2.38, 0.82, 0.82, RED)
    tf = o.text_frame; tf.vertical_anchor = MSO_ANCHOR.MIDDLE
    p = tf.paragraphs[0]; p.alignment = PP_ALIGN.CENTER
    run = p.add_run(); run.text = label
    _set_run(run, 9, True, False, WHITE)

# connector box
rect(sl, 0.82, 3.25, 3.4, 0.55, WHITE, MGRAY, 0.75)
txt(sl, 'Central editorial bottlenecks', 0.82, 3.25, 3.4, 0.55, 8, False, False, DARK, PP_ALIGN.CENTER)

rect(sl, 1.25, 3.9, 2.5, 0.38, LGRAY)
txt(sl, 'Web / Print separation', 1.25, 3.9, 2.5, 0.38, 8, False, False, GRAY, PP_ALIGN.CENTER)
rect(sl, 0.82, 4.38, 3.4, 0.38, LGRAY)
txt(sl, 'Multiple handovers & duplicate work', 0.82, 4.38, 3.4, 0.38, 8, False, False, GRAY, PP_ALIGN.CENTER)

# 6 bottom shared icons
icon_labels = ['Expertise', 'Editorial', 'SEO', 'AI Support', 'Data &\nInsights', 'Digital\nDistribution']
for k, label in enumerate(icon_labels):
    bx = 0.28 + k * 2.14
    oval(sl, bx, 5.12, 0.4, 0.4, LGRAY)
    rect(sl, bx - 0.8, 5.58, 2.0, 0.42, LGRAY)
    txt(sl, label, bx - 0.8, 5.58, 2.0, 0.42, 7.5, False, False, DARK, PP_ALIGN.CENTER)

# FUTURE DIRECTION column
shape_txt(sl, 'FUTURE DIRECTION', 7.1, 1.58, 6.1, 0.4, ORANGE, 11, True, WHITE, PP_ALIGN.CENTER)
txt(sl, 'Integrated consumer hubs', 7.1, 2.02, 6.1, 0.28, 8.5, False, True, GRAY, PP_ALIGN.CENTER)

# central consumer oval
co = oval(sl, 9.4, 2.55, 1.75, 1.75, DARK)
tf = co.text_frame; tf.vertical_anchor = MSO_ANCHOR.MIDDLE
p = tf.paragraphs[0]; p.alignment = PP_ALIGN.CENTER
run = p.add_run(); run.text = 'CONSUMER'
_set_run(run, 8, True, False, WHITE)
p2 = tf.add_paragraph(); p2.alignment = PP_ALIGN.CENTER
run2 = p2.add_run(); run2.text = 'NEEDS'
_set_run(run2, 8, True, False, WHITE)

# satellite hubs
hubs = [
    ('Personal\nFinance',    RED,    7.5,  2.1),
    ('High Tech\n& Mobility',ORANGE, 7.15, 3.5),
    ('Rights &\nProtection', ORANGE, 11.15,2.2),
    ('Health &\nNutrition',  ORANGE, 11.15,3.6),
    ('Housing &\nEnergy',    ORANGE, 9.2,  4.4),
]
for (label, col, hx, hy) in hubs:
    o = oval(sl, hx, hy, 1.3, 1.0, col)
    tf = o.text_frame; tf.vertical_anchor = MSO_ANCHOR.MIDDLE
    lines = label.split('\n')
    for j, line in enumerate(lines):
        if j == 0:
            p = tf.paragraphs[0]
        else:
            p = tf.add_paragraph()
        p.alignment = PP_ALIGN.CENTER
        run = p.add_run(); run.text = line
        _set_run(run, 7.5, True, False, WHITE)

# bottom benefits bar
rect(sl, 0.28, 6.22, 12.77, 0.52, RED)
benefits = ['✓  Faster collaboration', '✓  Less duplication', '✓  More scalability', '✓  Stronger consumer focus']
for k, b in enumerate(benefits):
    bx = 0.28 + k * (12.77 / 4)
    txt(sl, b, bx, 6.25, 12.77/4, 0.46, 9, True, False, WHITE, PP_ALIGN.CENTER)

logo(sl); pnum(sl, 3)


# ─────────────────────────────────────────────────────────────────────────────
# SLIDE 4 – "WHERE IS EFFICIENCY BEING LOST TODAY?"
# ─────────────────────────────────────────────────────────────────────────────
sl = prs.slides.add_slide(blank_layout)
rect(sl, 0, 0, 13.33, 7.5, WHITE)
badge(sl, 4, 0.52, 0.45)

mtxt(sl, [
    {'text': 'WHERE IS EFFICIENCY',  'size': 21, 'bold': True, 'color': ORANGE, 'align': PP_ALIGN.LEFT},
    {'text': 'BEING LOST TODAY?',    'size': 21, 'bold': True, 'color': RED,    'align': PP_ALIGN.LEFT},
], 0.9, 0.1, 10.5, 0.98)

divider(sl, 0.9, 1.1, 0.9)

rows = [
    ('◉', RED,    16, 'REPETITIVE CONSUMER REQUESTS',
     'High volume of Tier-1 questions handled manually (e.g., account, membership, basic info).'),
    ('⚙', ORANGE, 14, 'COMPLEX PRODUCTION WORKFLOWS',
     'Multiple handovers between experts, editorial, SEO, translation and design slow down publishing.'),
    ('◈', RED,    14, 'FR/NL DUPLICATION',
     'Full human translation and adaptation for many content pieces.'),
    ('●', ORANGE, 14, 'DUPLICATE DATA & REPORTING',
     'Siloed data and dashboards lead to duplication and inconsistent insights.'),
    ('◉', RED,    14, 'LOW-VALUE CONTENT VOLUME',
     'Resources spent on low-impact topics with limited consumer interest.'),
]

ry = 1.32
for i, (sym, scol, ssz, title, desc) in enumerate(rows):
    if i > 0:
        rect(sl, 0.28, ry, 8.3, 0.03, MGRAY)
    icon_oval(sl, sym, 0.28 + 0.325, ry + 0.08 + 0.325, 0.325, LGRAY, scol, ssz)
    txt(sl, title, 1.1, ry + 0.08, 7.4, 0.3, 9, True, False, ORANGE, PP_ALIGN.LEFT)
    txt(sl, desc,  1.1, ry + 0.42, 7.4, 0.55, 8.5, False, False, DARK, PP_ALIGN.LEFT)
    ry += 1.04

# opportunity box
rect(sl, 8.68, 1.25, 4.38, 5.95, OPPORT)
txt(sl, 'OPPORTUNITY', 8.68, 1.42, 4.38, 0.42, 12, True, False, RED, PP_ALIGN.CENTER)

# bullseye
cc = RGBColor(0xCC, 0x00, 0x33)
oval(sl, 9.54,  2.05, 2.7,  2.7,  cc)
oval(sl, 9.89,  2.40, 2.0,  2.0,  WHITE)
oval(sl, 10.24, 2.75, 1.3,  1.3,  RED)
oval(sl, 10.60, 3.11, 0.48, 0.48, WHITE)

txt(sl, 'Reduce friction,\nremove duplication\nand free up expertise\nfor higher-value work.',
    8.78, 4.95, 4.18, 1.85, 12, True, False, RED, PP_ALIGN.CENTER)

logo(sl); pnum(sl, 4)


# ─────────────────────────────────────────────────────────────────────────────
# SLIDE 5 – "WOW – OPPORTUNITIES FOR THE NEXT PHASE"
# ─────────────────────────────────────────────────────────────────────────────
sl = prs.slides.add_slide(blank_layout)
rect(sl, 0, 0, 13.33, 7.5, WHITE)
badge(sl, 5, 0.52, 0.37)

txt(sl, 'WOW – OPPORTUNITIES FOR THE NEXT PHASE',
    0.9, 0.1, 12.0, 0.5, 20, True, False, DARK, PP_ALIGN.LEFT)
divider(sl, 0.9, 0.65, 0.9)
txt(sl, 'Build on our strengths to create even more impact for consumers.',
    0.28, 0.73, 12.8, 0.32, 10, False, True, GRAY, PP_ALIGN.CENTER)

# column headers
shape_txt(sl, '1. TRUSTED CONSUMER\nPROTECTION PLATFORM', 0.28, 1.08, CW3, 0.62, RED,   9.5, True, WHITE, PP_ALIGN.CENTER)
shape_txt(sl, '2. AI-ENABLED DECISION\nSUPPORT',          4.57, 1.08, CW3, 0.62, ORANGE,9.5, True, WHITE, PP_ALIGN.CENTER)
shape_txt(sl, '3. LIFE MOMENT SUPPORT',                   8.86, 1.08, CW3, 0.62, AMBER, 9.5, True, WHITE, PP_ALIGN.CENTER)

# icon circles
oval(sl, 0.28 + CW3/2 - 0.32, 1.78, 0.64, 0.64, RED)
icon_oval(sl, '🛡', 0.28 + CW3/2, 1.78 + 0.32, 0.32, RED, WHITE, 13)

oval(sl, 4.57 + CW3/2 - 0.32, 1.78, 0.64, 0.64, ORANGE)
icon_oval(sl, '◉', 4.57 + CW3/2, 1.78 + 0.32, 0.32, ORANGE, WHITE, 13)

oval(sl, 8.86 + CW3/2 - 0.32, 1.78, 0.64, 0.64, AMBER)
icon_oval(sl, '⭐', 8.86 + CW3/2, 1.78 + 0.32, 0.32, AMBER, WHITE, 13)

# bullet columns
c1_items = ['Scams & fraud protection', 'Privacy & data', 'Dispute support', 'Contract optimisation', 'Stronger legal support']
c2_items = ['Smart comparators', 'Personalised recommendations', 'AI consumer assistant', 'Savings optimisation tools', 'Proactive alerts & guidance']
c3_items = ['Buying a home', 'Retirement', 'Investing', 'Energy transition', 'Major life decisions']
c_cols_all   = [c1_items, c2_items, c3_items]
c_dot_cols   = [RED, ORANGE, AMBER]
c_xs         = [0.28, 4.57, 8.86]

for ci in range(3):
    x = c_xs[ci]
    dot_col = c_dot_cols[ci]
    iy = 2.55
    for item in c_cols_all[ci]:
        oval(sl, x, iy + 0.06, 0.18, 0.18, dot_col)
        txt(sl, item, x + 0.24, iy, CW3 - 0.28, 0.28, 8.5, False, False, DARK, PP_ALIGN.LEFT)
        iy += 0.32

# callout boxes
callout(sl, 'Consumers expect more presence in Privacy/Data, Legal rights and Financial matters.\n(Brand Tracker Q4 2024)',
        0.28, 4.25, 4.17, 0.98, PINK_L, RED, RED, 7.5)
callout(sl, 'Leverage our expertise and data to deliver personalised, real-time consumer guidance.',
        4.57, 4.25, 4.17, 0.98, ORAN_L, ORANGE, ORANGE, 7.5)
callout(sl, 'Millennials show strong relevance for life moments and financial guidance.\n(Brand Tracker Q4 2024)',
        8.86, 4.25, 4.17, 0.98, YELL_L, AMBER, GOLD, 7.5)

# trophy banner
rect(sl, 0.28, 5.38, 12.77, 0.72, RED)
# trophy icon
o = oval(sl, 0.38, 5.44, 0.58, 0.58, AMBER)
tf = o.text_frame; tf.vertical_anchor = MSO_ANCHOR.MIDDLE
p = tf.paragraphs[0]; p.alignment = PP_ALIGN.CENTER
run = p.add_run(); run.text = '🏆'
_set_run(run, 14, True, False, WHITE)

txt(sl, 'Information built our reputation.   Helping consumers make better decisions should define our future.',
    1.1, 5.48, 11.7, 0.58, 11, True, False, WHITE, PP_ALIGN.CENTER)

logo(sl); pnum(sl, 5)


# ── save ─────────────────────────────────────────────────────────────────────
out = '/home/user/Sulli68/testachats_presentation.pptx'
prs.save(out)
print(f'Saved → {out}')
