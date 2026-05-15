#!/usr/bin/env python3
"""
Génération de la présentation testachats/testaankoop – 5 slides
"""
from pptx import Presentation
from pptx.util import Inches, Pt
from pptx.dml.color import RGBColor
from pptx.enum.text import PP_ALIGN, MSO_ANCHOR

prs = Presentation()
prs.slide_width  = Inches(13.33)
prs.slide_height = Inches(7.5)
blank = prs.slide_layouts[6]   # blank layout

# ── Palette ──────────────────────────────────────────────────────────────────
RED    = RGBColor(0xE8, 0x00, 0x3D)
ORANGE = RGBColor(0xF0, 0x80, 0x00)
AMBER  = RGBColor(0xF5, 0xA6, 0x23)
DARK   = RGBColor(0x1A, 0x1A, 0x1A)
GRAY   = RGBColor(0x77, 0x77, 0x77)
LGRAY  = RGBColor(0xF0, 0xF0, 0xF0)
WHITE  = RGBColor(0xFF, 0xFF, 0xFF)
PINK_L = RGBColor(0xFF, 0xE0, 0xEC)
ORAN_L = RGBColor(0xFF, 0xF0, 0xD8)
YELL_L = RGBColor(0xFF, 0xF8, 0xD0)
GOLD   = RGBColor(0x80, 0x60, 0x00)

def i(v):  return Inches(v)
def pt(v): return Pt(v)

# ── Primitive helpers ─────────────────────────────────────────────────────────
def _set_run(run, text, size, bold=False, italic=False, color=DARK):
    run.text = text
    run.font.size = pt(size)
    run.font.bold = bold
    run.font.italic = italic
    run.font.color.rgb = color

def _set_para(para, text, size, bold=False, italic=False, color=DARK,
              align=PP_ALIGN.LEFT, space_before=0):
    para.alignment = align
    if space_before:
        para.space_before = pt(space_before)
    if text:
        _set_run(para.add_run(), text, size, bold, italic, color)

# ── Shape (rectangle / oval) ──────────────────────────────────────────────────
def add_rect(slide, l, t, w, h, fill=None, line=None, lw=0.75):
    shp = slide.shapes.add_shape(1, i(l), i(t), i(w), i(h))
    if fill:
        shp.fill.solid()
        shp.fill.fore_color.rgb = fill
    else:
        shp.fill.background()
    if line:
        shp.line.color.rgb = line
        shp.line.width = pt(lw)
    else:
        shp.line.fill.background()
    return shp

def add_oval(slide, l, t, w, h, fill=RED):
    shp = slide.shapes.add_shape(9, i(l), i(t), i(w), i(h))
    shp.fill.solid()
    shp.fill.fore_color.rgb = fill
    shp.line.fill.background()
    return shp

# ── Shape with centred text (for headers, badges) ─────────────────────────────
def shape_txt(slide, text, l, t, w, h, fill=None, line=None, lw=0.75,
              size=10, bold=False, color=WHITE, align=PP_ALIGN.CENTER,
              italic=False, v=MSO_ANCHOR.MIDDLE):
    shp = add_rect(slide, l, t, w, h, fill=fill, line=line, lw=lw)
    tf  = shp.text_frame
    tf.word_wrap = True
    tf.vertical_anchor = v
    p = tf.paragraphs[0]
    p.alignment = align
    _set_run(p.add_run(), text, size, bold, italic, color)
    return shp

# ── Floating text box (single paragraph) ─────────────────────────────────────
def txt(slide, text, l, t, w, h, size=10, bold=False, color=DARK,
        align=PP_ALIGN.LEFT, italic=False):
    box = slide.shapes.add_textbox(i(l), i(t), i(w), i(h))
    tf  = box.text_frame
    tf.word_wrap = True
    p = tf.paragraphs[0]
    _set_para(p, text, size, bold, italic, color, align)
    return box

# ── Floating text box (multiple paragraphs) ───────────────────────────────────
def mtxt(slide, paras, l, t, w, h):
    """paras = list of dicts: text, size, bold, italic, color, align, space_before"""
    box = slide.shapes.add_textbox(i(l), i(t), i(w), i(h))
    tf  = box.text_frame
    tf.word_wrap = True
    for idx, pd in enumerate(paras):
        p = tf.paragraphs[0] if idx == 0 else tf.add_paragraph()
        _set_para(p,
                  pd.get('text', ''),
                  pd.get('size', 10),
                  pd.get('bold', False),
                  pd.get('italic', False),
                  pd.get('color', DARK),
                  pd.get('align', PP_ALIGN.LEFT),
                  pd.get('space_before', 0))
    return box

# ── Reusable slide furniture ──────────────────────────────────────────────────
def circle_num(slide, num, cx, cy, r=0.26, bg=RED):
    shp = add_oval(slide, cx - r, cy - r, r * 2, r * 2, fill=bg)
    tf  = shp.text_frame
    tf.vertical_anchor = MSO_ANCHOR.MIDDLE
    p = tf.paragraphs[0]
    p.alignment = PP_ALIGN.CENTER
    _set_run(p.add_run(), str(num), 14, bold=True, color=WHITE)

def divider(slide, l, t, w=0.8):
    add_rect(slide, l, t, w, 0.04, fill=RED)

def logo(slide, l=0.2, t=7.1):
    box = slide.shapes.add_textbox(i(l), i(t), i(1.5), i(0.35))
    tf  = box.text_frame
    p0  = tf.paragraphs[0]
    _set_run(p0.add_run(), 'testachats', 7, bold=True, color=RED)
    p1  = tf.add_paragraph()
    _set_run(p1.add_run(), 'testaankoop', 7, bold=True, color=DARK)

def page_num(slide, num):
    txt(slide, str(num), 12.9, 7.15, 0.3, 0.25, size=9, color=GRAY, align=PP_ALIGN.RIGHT)

def bullets(slide, items, cx, cy, cw, size=8.5, color=DARK):
    paras = [{'text': f'• {it}', 'size': size, 'color': color} for it in items]
    mtxt(slide, paras, cx, cy, cw, len(items) * 0.32 + 0.1)


# ═══════════════════════════════════════════════════════════════════════════════
# SLIDE 1 – What is truly essential for consumers?
# ═══════════════════════════════════════════════════════════════════════════════
s1 = prs.slides.add_slide(blank)
add_rect(s1, 0, 0, 13.33, 7.5, fill=WHITE)

circle_num(s1, 1, 0.5, 0.45)

mtxt(s1, [
    {'text': 'WHAT IS TRULY ESSENTIAL', 'size': 21, 'bold': True, 'color': DARK},
    {'text': 'FOR CONSUMERS?',          'size': 21, 'bold': True, 'color': RED},
], 0.85, 0.1, 10.0, 0.95)

divider(s1, 0.85, 1.1)
txt(s1, 'Our value is built on trust, expertise and impact.',
    0.25, 1.18, 10.5, 0.35, size=10.5)

# Four columns
CXS = [0.25, 3.5, 6.75, 10.0]
CW  = 3.0
CT  = 1.6
CH  = 3.9

col_headers = [
    ('TRUST &\nINDEPENDENCE',       DARK),
    ('PRACTICAL CONSUMER\nSUPPORT', ORANGE),
    ('EXPERTISE &\nDECISION SUPPORT', ORANGE),
    ('SIMPLICITY &\nGUIDANCE',      ORANGE),
]
col_bullets = [
    ['Independent voice', 'High credibility', 'Consumer-first mission'],
    ['Rights & disputes', 'Financial guidance', 'Energy', 'Major purchases', 'Contracts'],
    ['Testing capability', 'Structured data', 'Comparison intelligence', 'Help consumers make decisions'],
    ['Clear answers', 'Actionable advice', 'Save time & money', 'Build confidence'],
]

for k, cx in enumerate(CXS):
    add_rect(s1, cx, CT, CW, CH, fill=LGRAY)
    # icon placeholder
    add_oval(s1, cx + (CW - 0.55) / 2, CT + 0.1, 0.55, 0.55,
             fill=RGBColor(0xFF, 0xCC, 0xDD))
    hdr, hcol = col_headers[k]
    txt(s1, hdr, cx + 0.05, CT + 0.75, CW - 0.1, 0.55,
        size=8.5, bold=True, color=hcol, align=PP_ALIGN.CENTER)
    bullets(s1, col_bullets[k], cx + 0.15, CT + 1.4, CW - 0.25)

# ── Stats ──
add_rect(s1, 0.25, 5.62, 2.95, 1.65, fill=LGRAY)
txt(s1, '96%', 0.25, 5.67, 2.95, 0.75,
    size=34, bold=True, color=DARK, align=PP_ALIGN.CENTER)
txt(s1, 'aided awareness\n(Brand Tracker Q4 2024)',
    0.25, 6.43, 2.95, 0.5, size=7.5, color=GRAY, align=PP_ALIGN.CENTER)

txt(s1, 'Top relevance scores:', 3.25, 5.58, 3.0, 0.3, size=8, bold=True, color=ORANGE)
mtxt(s1, [
    {'text': '78% fiscal questions',      'size': 8, 'color': DARK},
    {'text': '76% health questions',      'size': 8, 'color': DARK},
    {'text': '72% property decisions',    'size': 8, 'color': DARK},
    {'text': '(Brand Tracker Q4 2024)',   'size': 7, 'color': GRAY, 'italic': True},
], 3.25, 5.88, 3.1, 0.95)

txt(s1, 'Top traffic hubs:', 6.5, 5.58, 3.0, 0.3, size=8, bold=True, color=ORANGE)
mtxt(s1, [
    {'text': '20% Personal Finance',  'size': 8, 'color': DARK},
    {'text': '16.8% Home Equipment',  'size': 8, 'color': DARK},
    {'text': '16.8% High Tech',       'size': 8, 'color': DARK},
    {'text': '(Hub Online Sessions)', 'size': 7, 'color': GRAY, 'italic': True},
], 6.5, 5.88, 3.1, 0.95)

add_rect(s1, 9.75, 5.58, 3.35, 1.7, fill=PINK_L)
txt(s1, 'Relevance is the #1 driver of consideration (43%)\n(Brand Tracker Q4 2024)',
    9.85, 5.65, 3.15, 1.5, size=8.5, color=RED, italic=True)

page_num(s1, 1)
logo(s1)

# ═══════════════════════════════════════════════════════════════════════════════
# SLIDE 2 – Where should we prioritise, simplify or stop?
# ═══════════════════════════════════════════════════════════════════════════════
s2 = prs.slides.add_slide(blank)
add_rect(s2, 0, 0, 13.33, 7.5, fill=WHITE)

circle_num(s2, 2, 0.5, 0.45)

mtxt(s2, [
    {'text': 'WHERE SHOULD WE PRIORITISE,', 'size': 20, 'bold': True, 'color': DARK},
    {'text': 'SIMPLIFY OR STOP?',            'size': 20, 'bold': True, 'color': RED},
], 0.85, 0.1, 11.5, 0.95)

divider(s2, 0.85, 1.08)
txt(s2, 'Focus on what creates the most value. Simplify to increase impact.',
    0.25, 1.15, 12.8, 0.3, size=10, align=PP_ALIGN.CENTER)

C1, C2, C3, CW2 = 0.25, 4.65, 9.05, 4.2

shape_txt(s2, 'PRIORITISE MORE', C1, 1.53, CW2, 0.42, fill=RED,
          size=11, bold=True, color=WHITE, align=PP_ALIGN.CENTER)
shape_txt(s2, 'SIMPLIFY', C2, 1.53, CW2, 0.42, fill=ORANGE,
          size=11, bold=True, color=WHITE, align=PP_ALIGN.CENTER)
shape_txt(s2, 'REDUCE / REASSESS', C3, 1.53, CW2, 0.42, fill=AMBER,
          size=11, bold=True, color=WHITE, align=PP_ALIGN.CENTER)

txt(s2, 'Invest in high-value domains',     C1+0.1, 2.0, CW2-0.1, 0.28, size=8.5, bold=True)
txt(s2, 'Reduce operational complexity',    C2+0.1, 2.0, CW2-0.1, 0.28, size=8.5, bold=True)
txt(s2, 'Rationalise lower impact areas',   C3+0.1, 2.0, CW2-0.1, 0.28, size=8.5, bold=True)

items_c1 = ['Personal Finance', 'Rights & Protection', 'Housing & Energy',
            'High Tech', 'High-impact consumer decisions']
items_c2 = ['Shared hub-based workflows', 'Reduce duplicate layers',
            'Simplify web/print production', 'AI-assisted repetitive tasks',
            'Faster multi-channel publishing']
items_c3 = ['Low-interest niche testing', 'Fragmented micro-topics',
            'Duplicate workflows', 'Separation between publication brands',
            'Heavy manual adaptation processes']

for items, cx2 in [(items_c1, C1), (items_c2, C2), (items_c3, C3)]:
    bullets(s2, items, cx2 + 0.1, 2.33, CW2 - 0.15)

# callout boxes
add_rect(s2, C1, 4.28, CW2, 0.92, fill=PINK_L, line=RED, lw=0.5)
txt(s2, '54% of online sessions concentrated in Personal Finance, '
        'Home Equipment and High Tech.\n(Hub Online Sessions)',
    C1+0.12, 4.33, CW2-0.2, 0.82, size=7.5, color=RED, italic=True)

add_rect(s2, C2, 4.28, CW2, 0.92, fill=ORAN_L, line=ORANGE, lw=0.5)
txt(s2, 'Major opportunity to reduce complexity and speed up publishing.\n(Internal observation)',
    C2+0.12, 4.33, CW2-0.2, 0.82, size=7.5, color=ORANGE, italic=True)

add_rect(s2, C3, 4.28, CW2, 0.92, fill=YELL_L, line=AMBER, lw=0.5)
txt(s2, 'Lower traffic hubs:\nShopping 3.1% | Rights 5.4% | Mobility 5.5%\n(Hub Online Sessions)',
    C3+0.12, 4.33, CW2-0.2, 0.82, size=7.5, color=GOLD, italic=True)

# AI / Human section
add_rect(s2, 0.25, 5.3, 6.2, 1.9, fill=LGRAY)
txt(s2, 'AI SHOULD SUPPORT', 0.35, 5.38, 6.0, 0.3, size=8.5, bold=True)
mtxt(s2, [
    {'text': '• First drafts',           'size': 8.5, 'color': DARK},
    {'text': '• Translation & SEO',      'size': 8.5, 'color': DARK},
    {'text': '• Workflow acceleration',  'size': 8.5, 'color': DARK},
], 0.45, 5.72, 3.0, 0.75)

add_rect(s2, 6.58, 5.3, 6.5, 1.9, fill=LGRAY)
txt(s2, 'HUMAN EXPERTISE REMAINS ESSENTIAL FOR',
    6.68, 5.38, 6.2, 0.3, size=8.5, bold=True)
mtxt(s2, [
    {'text': '• Legal interpretation',     'size': 8.5, 'color': DARK},
    {'text': '• Validation & credibility', 'size': 8.5, 'color': DARK},
    {'text': '• Consumer advocacy',        'size': 8.5, 'color': DARK},
], 6.68, 5.72, 3.5, 0.75)

page_num(s2, 2)
logo(s2)

# ═══════════════════════════════════════════════════════════════════════════════
# SLIDE 3 – Operational Consequences & Future Organisation
# ═══════════════════════════════════════════════════════════════════════════════
s3 = prs.slides.add_slide(blank)
add_rect(s3, 0, 0, 13.33, 7.5, fill=WHITE)

circle_num(s3, 3, 0.5, 0.45)

mtxt(s3, [
    {'text': 'OPERATIONAL CONSEQUENCES',  'size': 20, 'bold': True, 'color': DARK},
    {'text': '& FUTURE ORGANISATION',     'size': 20, 'bold': True, 'color': RED},
], 0.85, 0.1, 11.5, 0.95)

divider(s3, 0.85, 1.08)
txt(s3, 'From publication silos to integrated consumer hubs.',
    0.25, 1.15, 12.0, 0.3, size=10)

# TODAY column
shape_txt(s3, 'TODAY', 0.25, 1.53, 6.1, 0.38, fill=LGRAY,
          size=11, bold=True, color=DARK, align=PP_ALIGN.CENTER)
txt(s3, 'Publication-centric & fragmented',
    0.25, 1.95, 6.1, 0.28, size=8.5, color=GRAY, align=PP_ALIGN.CENTER, italic=True)

# TA / BD / TS / TAI bubbles
for label, bx in [('TA', 0.6), ('BD', 1.85), ('TS', 3.1), ('TAI', 4.35)]:
    shp = add_oval(s3, bx, 2.3, 0.8, 0.8, fill=RED)
    tf  = shp.text_frame
    tf.vertical_anchor = MSO_ANCHOR.MIDDLE
    p   = tf.paragraphs[0]
    p.alignment = PP_ALIGN.CENTER
    _set_run(p.add_run(), label, 9, bold=True, color=WHITE)

add_rect(s3, 1.0, 3.25, 3.3, 0.55, fill=WHITE, line=GRAY, lw=0.75)
txt(s3, 'Central editorial bottlenecks', 1.0, 3.3, 3.3, 0.45,
    size=8, align=PP_ALIGN.CENTER)
txt(s3, 'Web / Print separation',
    1.2, 3.9, 2.8, 0.35, size=8, color=GRAY, align=PP_ALIGN.CENTER)
txt(s3, 'Multiple handovers & duplicate work',
    1.0, 4.35, 3.4, 0.35, size=8, color=GRAY, align=PP_ALIGN.CENTER)

# FUTURE DIRECTION column
shape_txt(s3, 'FUTURE DIRECTION', 7.0, 1.53, 6.1, 0.38, fill=ORANGE,
          size=11, bold=True, color=WHITE, align=PP_ALIGN.CENTER)
txt(s3, 'Integrated consumer hubs',
    7.0, 1.95, 6.1, 0.28, size=8.5, color=GRAY, align=PP_ALIGN.CENTER, italic=True)

# Consumer NEEDS centre
shp_cn = add_oval(s3, 9.3, 2.5, 1.8, 1.8, fill=DARK)
tf_cn  = shp_cn.text_frame
tf_cn.vertical_anchor = MSO_ANCHOR.MIDDLE
p_cn   = tf_cn.paragraphs[0]
p_cn.alignment = PP_ALIGN.CENTER
_set_run(p_cn.add_run(), 'CONSUMER\nNEEDS', 8, bold=True, color=WHITE)

# Satellite hubs
future_hubs = [
    ('Personal\nFinance',   7.5,  2.1, RED),
    ('High Tech\n& Mobility', 7.1, 3.5, ORANGE),
    ('Rights &\nProtection', 11.0, 2.2, ORANGE),
    ('Health &\nNutrition',  11.0, 3.6, ORANGE),
    ('Housing &\nEnergy',    9.1,  4.35, ORANGE),
]
for label, hx, hy, hc in future_hubs:
    shp_h = add_oval(s3, hx, hy, 1.3, 1.0, fill=hc)
    tf_h  = shp_h.text_frame
    tf_h.vertical_anchor = MSO_ANCHOR.MIDDLE
    p_h   = tf_h.paragraphs[0]
    p_h.alignment = PP_ALIGN.CENTER
    _set_run(p_h.add_run(), label, 7.5, bold=True, color=WHITE)

# Icon row
icons_s3 = ['Expertise', 'Editorial', 'SEO', 'AI Support', 'Data & Insights', 'Digital Distribution']
for ki, name in enumerate(icons_s3):
    shape_txt(s3, name, 0.25 + ki * 2.14, 5.55, 2.0, 0.52, fill=LGRAY,
              size=7.5, bold=False, color=DARK, align=PP_ALIGN.CENTER)

# Bottom bar
add_rect(s3, 0.25, 6.18, 12.85, 0.52, fill=RED)
for kb, benefit in enumerate(['Faster collaboration', 'Less duplication',
                                'More scalability', 'Stronger consumer focus']):
    txt(s3, f'✓  {benefit}', 0.5 + kb * 3.2, 6.22, 3.0, 0.42,
        size=9, bold=True, color=WHITE, align=PP_ALIGN.CENTER)

page_num(s3, 3)
logo(s3)

# ═══════════════════════════════════════════════════════════════════════════════
# SLIDE 4 – Where is efficiency being lost today?
# ═══════════════════════════════════════════════════════════════════════════════
s4 = prs.slides.add_slide(blank)
add_rect(s4, 0, 0, 13.33, 7.5, fill=WHITE)

circle_num(s4, 4, 0.5, 0.45)

mtxt(s4, [
    {'text': 'WHERE IS EFFICIENCY',  'size': 20, 'bold': True, 'color': ORANGE},
    {'text': 'BEING LOST TODAY?',    'size': 20, 'bold': True, 'color': RED},
], 0.85, 0.1, 10.0, 0.95)

divider(s4, 0.85, 1.08)

eff_data = [
    ('REPETITIVE CONSUMER REQUESTS',
     'High volume of Tier-1 questions handled manually (e.g., account, membership, basic info).'),
    ('COMPLEX PRODUCTION WORKFLOWS',
     'Multiple handovers between experts, editorial, SEO, translation and design slow down publishing.'),
    ('FR/NL DUPLICATION',
     'Full human translation and adaptation for many content pieces.'),
    ('DUPLICATE DATA & REPORTING',
     'Siloed data and dashboards lead to duplication and inconsistent insights.'),
    ('LOW-VALUE CONTENT VOLUME',
     'Resources spent on low-impact topics with limited consumer interest.'),
]

for ki, (title, desc) in enumerate(eff_data):
    ry = 1.35 + ki * 1.05
    add_oval(s4, 0.25, ry + 0.05, 0.65, 0.65, fill=LGRAY)
    txt(s4, title, 1.05, ry + 0.05, 7.1, 0.3, size=9, bold=True, color=ORANGE)
    txt(s4, desc,  1.05, ry + 0.38, 7.1, 0.6, size=8.5)

# Opportunity box
add_rect(s4, 8.5, 1.25, 4.6, 5.8, fill=RGBColor(0xFF, 0xCC, 0xD8))
txt(s4, 'OPPORTUNITY', 8.5, 1.38, 4.6, 0.42,
    size=12, bold=True, color=RED, align=PP_ALIGN.CENTER)

shp_t = add_oval(s4, 9.55, 2.05, 2.5, 2.5, fill=RGBColor(0xCC, 0x00, 0x33))
tf_t  = shp_t.text_frame
tf_t.vertical_anchor = MSO_ANCHOR.MIDDLE
p_t   = tf_t.paragraphs[0]
p_t.alignment = PP_ALIGN.CENTER
_set_run(p_t.add_run(), '◎', 50, color=WHITE)   # bullseye char

txt(s4, 'Reduce friction,\nremove duplication\nand free up expertise\nfor higher-value work.',
    8.6, 4.7, 4.3, 1.7, size=12, bold=True, color=RED, align=PP_ALIGN.CENTER)

page_num(s4, 4)
logo(s4)

# ═══════════════════════════════════════════════════════════════════════════════
# SLIDE 5 – WOW – Opportunities for the next phase
# ═══════════════════════════════════════════════════════════════════════════════
s5 = prs.slides.add_slide(blank)
add_rect(s5, 0, 0, 13.33, 7.5, fill=WHITE)

circle_num(s5, 5, 0.5, 0.37)
txt(s5, 'WOW – OPPORTUNITIES FOR THE NEXT PHASE',
    0.85, 0.1, 12.0, 0.48, size=20, bold=True)
divider(s5, 0.85, 0.63)
txt(s5, 'Build on our strengths to create even more impact for consumers.',
    0.25, 0.72, 12.8, 0.3, size=10, align=PP_ALIGN.CENTER)

C1_5, C2_5, C3_5, CW5 = 0.25, 4.65, 9.05, 4.2

shape_txt(s5, '1. TRUSTED CONSUMER\nPROTECTION PLATFORM', C1_5, 1.12, CW5, 0.58,
          fill=RED, size=9.5, bold=True, color=WHITE, align=PP_ALIGN.CENTER)
shape_txt(s5, '2. AI-ENABLED DECISION\nSUPPORT', C2_5, 1.12, CW5, 0.58,
          fill=ORANGE, size=9.5, bold=True, color=WHITE, align=PP_ALIGN.CENTER)
shape_txt(s5, '3. LIFE MOMENT SUPPORT', C3_5, 1.12, CW5, 0.58,
          fill=AMBER, size=9.5, bold=True, color=WHITE, align=PP_ALIGN.CENTER)

# small icon circles
for cx5, cf in [(C1_5 + CW5/2 - 0.35, RED),
                (C2_5 + CW5/2 - 0.35, ORANGE),
                (C3_5 + CW5/2 - 0.35, AMBER)]:
    add_oval(s5, cx5, 1.78, 0.7, 0.7, fill=cf)

col5_1 = ['Scams & fraud protection', 'Privacy & data', 'Dispute support',
           'Contract optimisation', 'Stronger legal support']
col5_2 = ['Smart comparators', 'Personalised recommendations',
           'AI consumer assistant', 'Savings optimisation tools',
           'Proactive alerts & guidance']
col5_3 = ['Buying a home', 'Retirement', 'Investing',
           'Energy transition', 'Major life decisions']

for col_items, cx5 in [(col5_1, C1_5), (col5_2, C2_5), (col5_3, C3_5)]:
    bullets(s5, col_items, cx5 + 0.1, 2.6, CW5 - 0.15)

add_rect(s5, C1_5, 4.32, CW5, 0.95, fill=PINK_L, line=RED, lw=0.5)
txt(s5, 'Consumers expect more presence in Privacy/Data, Legal rights and '
        'Financial matters.\n(Brand Tracker Q4 2024)',
    C1_5+0.12, 4.37, CW5-0.2, 0.85, size=7.5, color=RED, italic=True)

add_rect(s5, C2_5, 4.32, CW5, 0.95, fill=ORAN_L, line=ORANGE, lw=0.5)
txt(s5, 'Leverage our expertise and data to deliver personalised, real-time consumer guidance.',
    C2_5+0.12, 4.37, CW5-0.2, 0.85, size=7.5, color=ORANGE, italic=True)

add_rect(s5, C3_5, 4.32, CW5, 0.95, fill=YELL_L, line=AMBER, lw=0.5)
txt(s5, 'Millennials show strong relevance for life moments and financial guidance.\n(Brand Tracker Q4 2024)',
    C3_5+0.12, 4.37, CW5-0.2, 0.85, size=7.5, color=GOLD, italic=True)

# Bottom trophy banner
add_rect(s5, 0.25, 5.42, 12.85, 0.72, fill=RED)
add_oval(s5, 0.35, 5.48, 0.6, 0.6, fill=AMBER)
txt(s5, 'Information built our reputation.   '
        'Helping consumers make better decisions should define our future.',
    1.1, 5.5, 11.7, 0.62, size=10.5, bold=True, color=WHITE, align=PP_ALIGN.CENTER)

page_num(s5, 5)
logo(s5)

# ── Save ──────────────────────────────────────────────────────────────────────
out = '/home/user/Sulli68/testachats_presentation.pptx'
prs.save(out)
print(f'Saved → {out}')
